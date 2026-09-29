# -*- coding: utf-8 -*-
"""add_scenarios_a93_fp24_2026_09_30.py - close the JA-116 coverage gap for A93 and FP-24.

WHY. JA-116 requires every A-NN audit category and FP-NN false-positive class in
prompts/Japanese language Accuracy check.txt to be named in at least one scenario row of
specifications/test-scenarios-by-specialist-perspective.xlsx. A93 (MEANING-GLOSS SENSE-COVERAGE vs
EXAMPLES) and FP-24 (polysemous-surface substring does not require a multi-sense gloss) were both
added on 2026-06-09 with the F.60 propagation, and their scenario rows were never written. The
content-integrity workflow has been red on master ever since - it is the exact prompts-to-xlsx
drift JA-116 exists to catch, and it caught it.

WHY NOT tools/sync_test_scenarios_with_prompts_feedback_2026_05_17.py, which JA-116's message
names. That script generated the 2026-05-17 bulk sync, and its output shape (A-147/A-148/A-149,
FP-20/FP-21) is sparser than what the project has written by hand since: A-154 and A-155 carry a
real Severity, Test type, Notes, Effort, Owner and Tools, and put the FP code in the Scenario text
under an `A-NNN` id. Regenerating would either not match that convention or rewrite rows nobody
asked to touch. Two rows appended in the current house style is the smaller, more faithful change.

Both rows follow A-155 exactly: same column order, same Sub-category vocabulary, same
"Reference: prompts/... <CODE>" first test step, same bounded Expected-result phrasing that
Rule 4's writing discipline requires ("against the current corpus snapshot", not "0 findings").

Run:  python tools/add_scenarios_a93_fp24_2026_09_30.py
"""
import os
import shutil
import sys
from datetime import datetime

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
import openpyxl

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
XLSX = os.path.join(ROOT, "specifications", "test-scenarios-by-specialist-perspective.xlsx")
TAB = "A. Japanese language"

# ID, Sub-category, Persp #, Scenario, Test steps, Expected result, Priority, Severity if fails,
# Test type, Notes, Estimated effort, Owner / role, Tools / scripts required
ROWS = [
    ("A-156", "Accuracy prompt category", "1",
     "Accuracy prompt audit category A93: MEANING-GLOSS SENSE-COVERAGE vs EXAMPLES - a pattern's "
     "meaning_en must be the union of exactly the senses its examples exercise, in both directions",
     "1. Reference: prompts/Japanese language Accuracy check.txt A93. "
     "2. For each pattern, segment meaning_en into sense-clauses and map every example to a "
     "covering clause. "
     "3. TOO NARROW: an example exercising a sense no clause covers (って glossed only "
     "\"named; called\" while an example is casual hearsay). "
     "4. OVERREACH: a clause no shown example demonstrates (the cleft-focus のは〜だ \"reason "
     "for\" case). "
     "5. Prioritise polysemous particles and enders: って, と, し, から, のに, ば/たら.",
     "Every example maps to a gloss clause and every gloss clause has a demonstrating example, "
     "against the current corpus snapshot; residual cases documented as FP-24 substring "
     "coincidences rather than defects.",
     "P3", "Major", "Manual",
     "Trigger: native review of the standalone N4 grammar deliverable, rounds 5-6. Direction A "
     "(gloss clause with no example) is machine-checkable; direction B (example with no clause) is "
     "native judgment - see FP-24 and row A-157. Added 2026-09-30 to close the JA-116 gap that had "
     "left content-integrity red on master since 2026-06-09.",
     "1h", "Native-teacher reviewer",
     "prompts/Japanese language Accuracy check.txt A93; tools/check_content_integrity.py JA-116"),

    ("A-157", "False-positive class", "1",
     "FP-24 polysemous-surface substring does NOT require a multi-sense gloss - a single-sense "
     "gloss is not a defect merely because the pattern's surface form contains a polysemous "
     "particle",
     "1. Reference: prompts/Japanese language Accuracy check.txt FP-24. "
     "2. Before filing any A93 too-narrow finding, check whether the polysemous particle appears "
     "as a substring of the pattern name rather than as a functioning particle. "
     "3. Confirm against the DIR-B run that returned 58 of 178 N5 patterns, all correctly "
     "mono-sense: しか〜ない, など, 何, だれ, から〜まで. "
     "4. Do not auto-flag on surface-form grounds; route direction B to a native reviewer.",
     "No A93 too-narrow finding is raised on surface-form grounds alone against the current "
     "corpus snapshot; DIR-B output is treated as a candidate list for native review, not as a "
     "defect list.",
     "P3", "Medium", "Documentation + audit",
     "Added 2026-06-09 during F.60 DIR-B validation; the scenario row was missed in that commit. "
     "Pairs with row A-156. The class exists because an automated too-narrow check floods: 58/178 "
     "matched on substring coincidence.",
     "30m", "Auditor",
     "prompts/Japanese language Accuracy check.txt FP-24; tools/check_content_integrity.py JA-116"),
]


def main() -> int:
    wb = openpyxl.load_workbook(XLSX)
    sh = wb[TAB]
    before_rows = sh.max_row
    existing = {str(sh.cell(r, 1).value) for r in range(5, before_rows + 1)}
    for r in ROWS:
        assert r[0] not in existing, "%s already present" % r[0]

    shutil.copy2(XLSX, os.path.join(
        ROOT, "specifications",
        "test-scenarios-by-specialist-perspective.bak_ja116_%s.xlsx"
        % datetime.now().strftime("%Y%m%d_%H%M%S")))

    # Mirror the formatting of the last data row so the appended rows are not style islands.
    template = before_rows
    for i, row in enumerate(ROWS):
        dest = before_rows + 1 + i
        for col, value in enumerate(row, start=1):
            c = sh.cell(dest, col, value)
            src = sh.cell(template, col)
            c.font = src.font.copy()
            c.alignment = src.alignment.copy()
            c.border = src.border.copy()
            c.fill = src.fill.copy()
        sh.row_dimensions[dest].height = sh.row_dimensions[template].height
    wb.save(XLSX)

    # ---- verify from a fresh read
    wb2 = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    sh2 = wb2[TAB]
    got = list(sh2.iter_rows(min_row=5, values_only=True))
    ids = [str(r[0]) for r in got]
    assert ids[-2:] == [r[0] for r in ROWS], "appended rows not found at the tail: %s" % ids[-3:]
    print("%s: %d rows -> %d rows" % (TAB, before_rows - 4, len(got)))
    for r in ROWS:
        print("   + %-6s %-26s %s" % (r[0], r[1], r[3][:72]))
    print("\nsheets in workbook: %d (unchanged: %s)"
          % (len(wb2.sheetnames), len(wb2.sheetnames) == len(wb.sheetnames)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
